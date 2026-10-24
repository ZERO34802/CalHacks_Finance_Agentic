"""Versioned deterministic adapters: CSV / XLSX / API JSON -> Invoice. Owner: A.

Code recomputes every line amount (qty x unit price). A printed line amount that differs
becomes an invoice flag `line_amount_mismatch:<line_no>:<printed>` (Matcher turns it into
total_mismatch). Cells starting with = + - @ are neutralized and flagged; in a money field
they raise IngestError (quarantine), unless the cell is a plain number like "-12.50".
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date
from decimal import Decimal
from typing import Any, Optional

from backend.pipeline.lib.money import parse_decimal, q2
from contracts import Invoice, LineItem

ADAPTER_VERSIONS = {"csv": "csv.invoice.v1", "xlsx": "xlsx.invoice.v1",
                    "kestrel.invoice.v2": "kestrel.invoice.v2", "aerolane.invoice.v1": "aerolane.invoice.v1"}

COLUMNS = ["invoice_number", "invoice_date", "po_number", "currency", "line_no", "sku", "description",
           "quantity", "unit_price", "line_amount", "remit_last4"]
HEADER_FIELDS = ["invoice_number", "invoice_date", "po_number", "currency", "remit_last4"]
MONEY_FIELDS = {"quantity", "unit_price", "line_amount", "total", "qty", "rate", "amount", "invoice_total"}
_INJECT = ("=", "+", "-", "@")
_ISO_CCY = re.compile(r"^[A-Z]{3}$")


class IngestError(ValueError):
    """Payload fails validation -> quarantine with this reason."""


def _neutralize(value: Any, field: str, flags: list[str]) -> Any:
    if not isinstance(value, str):
        return value
    v = value.strip()
    if v and v[0] in _INJECT:
        if field in MONEY_FIELDS:
            try:
                parse_decimal(v, field)
                return v  # a plain signed number is data, not a formula
            except ValueError:
                raise IngestError(f"formula injection in money field {field!r}: {v[:40]!r}")
        flags.append(f"formula_injection:{field}")
        return "'" + v  # neutralized for display/export
    return v


def _date(v: Any, field: str) -> date:
    try:
        d = date.fromisoformat(str(v).strip())
    except ValueError:
        raise IngestError(f"{field}: not an ISO date: {v!r}")
    if not (2000 <= d.year <= 2100):
        raise IngestError(f"{field}: implausible date {d}")
    return d


def _ccy(v: Any) -> str:
    s = str(v or "").strip().upper()
    if not _ISO_CCY.match(s):
        raise IngestError(f"currency: not an ISO code: {v!r}")
    return s


def _num(v: Any, field: str) -> Decimal:
    try:
        return parse_decimal(v, field)
    except ValueError as e:
        raise IngestError(str(e))


def _last4(v: Any) -> Optional[str]:
    s = re.sub(r"\D", "", str(v or ""))
    return s[-4:] if s else None


def _line(n: int, sku: str, desc: str, qty: Any, price: Any, printed: Any, flags: list[str]) -> LineItem:
    if not sku:
        raise IngestError(f"line {n}: missing sku")
    q, p = _num(qty, "quantity"), _num(price, "unit_price")
    amount = q2(q * p)
    if printed not in (None, ""):
        pa = _num(printed, "line_amount")
        if q2(pa) != amount:
            flags.append(f"line_amount_mismatch:{n}:{q2(pa)}")
    return LineItem(line_no=n, sku=sku, description=desc or "", qty=q, unit_price=p, amount=amount)


def _invoice(doc_id: str, vendor_id: str, inv_no: Any, inv_date: Any, ccy: Any, po: Any, lines: list[LineItem],
             total: Optional[Decimal], bank: Any, flags: list[str]) -> Invoice:
    if not str(inv_no or "").strip():
        raise IngestError("missing invoice number")
    if not lines:
        raise IngestError("invoice has no lines")
    computed = q2(sum((li.amount for li in lines), Decimal(0)))
    return Invoice(doc_id=doc_id, vendor_id=vendor_id, invoice_no=str(inv_no).strip(),
                   invoice_date=_date(inv_date, "invoice_date"), currency=_ccy(ccy),
                   po_number=(str(po).strip() or None) if po else None, lines=lines,
                   total=q2(total) if total is not None else computed, bank_last4=_last4(bank), flags=flags)


def _from_rows(rows: list[dict], doc_id: str, vendor_id: str, flags: list[str]) -> Invoice:
    if not rows:
        raise IngestError("no data rows")
    missing = [c for c in COLUMNS if c not in rows[0]]
    if missing:
        raise IngestError(f"unknown layout, missing columns {missing}")
    clean = [{k: _neutralize(v, k, flags) for k, v in r.items() if k in COLUMNS} for r in rows]
    head = clean[0]
    for i, r in enumerate(clean[1:], 2):
        for f in HEADER_FIELDS:
            if str(r[f]).strip() != str(head[f]).strip():
                raise IngestError(f"row {i}: header field {f} differs between rows")
    lines = []
    for i, r in enumerate(clean, 1):
        try:
            n = int(str(r["line_no"]).strip() or i)
        except ValueError:
            raise IngestError(f"row {i}: bad line_no {r['line_no']!r}")
        lines.append(_line(n, str(r["sku"]).strip(), str(r["description"] or "").strip(),
                           r["quantity"], r["unit_price"], r["line_amount"], flags))
    printed = q2(sum((_num(r["line_amount"], "line_amount") for r in clean), Decimal(0)))
    return _invoice(doc_id, vendor_id, head["invoice_number"], head["invoice_date"], head["currency"],
                    head["po_number"], lines, printed, head["remit_last4"], flags)


def parse_csv(data: bytes, doc_id: str, vendor_id: str) -> Invoice:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise IngestError("CSV is not UTF-8")
    first = text.splitlines()[0] if text else ""
    delim = ";" if first.count(";") > first.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    reader.fieldnames = [f.strip().lower() for f in (reader.fieldnames or [])]
    rows = [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]
    return _from_rows(rows, doc_id, vendor_id, [])


def parse_xlsx(data: bytes, doc_id: str, vendor_id: str) -> Invoice:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as e:
        raise IngestError(f"unreadable xlsx: {type(e).__name__}")
    if "Invoice" not in wb.sheetnames:
        raise IngestError(f"unknown layout: no 'Invoice' sheet (sheets: {wb.sheetnames})")
    it = wb["Invoice"].iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(it, [])]
    rows = []
    for vals in it:
        if vals is None or all(v is None or str(v).strip() == "" for v in vals):
            continue
        rows.append({h: ("" if v is None else (v.isoformat()[:10] if hasattr(v, "isoformat") else
                                               str(v)))
                     for h, v in zip(header, vals) if h})
    return _from_rows(rows, doc_id, vendor_id, [])


def _req(p: dict, key: str) -> Any:
    if key not in p or p[key] in (None, ""):
        raise IngestError(f"missing required field {key!r}")
    return p[key]


def parse_api(data: bytes, doc_id: str, vendor_id: str) -> tuple[Invoice, str]:
    try:
        p = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as e:
        raise IngestError(f"invalid JSON: {e}")
    if not isinstance(p, dict):
        raise IngestError("payload is not a JSON object")
    schema = p.get("schema")
    flags: list[str] = []
    p = {k: _neutralize(v, k, flags) for k, v in p.items()}
    if schema == "kestrel.invoice.v2":
        known = {"schema", "invoice_number", "invoice_date", "po_number", "currency", "remit_account_last4",
                 "lines", "total", "due_date"}
        items = _req(p, "lines")
        mk = lambda i, li: _line(i, str(_req(li, "sku")).strip(), str(li.get("description", "")),  # noqa: E731
                                 _req(li, "quantity"), _req(li, "unit_price"), li.get("amount"), flags)
        head = ("invoice_number", "invoice_date", "currency", "po_number", "total", "remit_account_last4")
    elif schema == "aerolane.invoice.v1":
        known = {"schema", "inv_no", "date", "po", "ccy", "bank_last4", "items", "invoice_total"}
        items = _req(p, "items")
        mk = lambda i, li: _line(i, str(_req(li, "code")).strip(), str(li.get("text", "")),  # noqa: E731
                                 _req(li, "qty"), _req(li, "rate"), li.get("amount"), flags)
        head = ("inv_no", "date", "ccy", "po", "invoice_total", "bank_last4")
    else:
        raise IngestError(f"unknown API schema {schema!r}")
    extra = set(p) - known
    if extra:  # format drift or added payment fields: never auto-mapped
        raise IngestError(f"unknown fields for {schema}: {sorted(extra)}")
    if not isinstance(items, list) or not all(isinstance(li, dict) for li in items):
        raise IngestError("lines must be a list of objects")
    lines = [mk(i, {k: _neutralize(v, k, flags) for k, v in li.items()}) for i, li in enumerate(items, 1)]
    no, dt, ccy, po, tot, bank = head
    inv = _invoice(doc_id, vendor_id, _req(p, no), _req(p, dt), _req(p, ccy), p.get(po), lines,
                   _num(_req(p, tot), "total"), p.get(bank), flags)
    if schema == "kestrel.invoice.v2" and p.get("due_date"):
        inv.due_date = _date(p["due_date"], "due_date")
    return inv, schema


def parse_structured(kind: str, data: bytes, doc_id: str, vendor_id: str) -> tuple[Invoice, str]:
    """kind is the sniffed kind: csv | xlsx | json. Returns (invoice, adapter version)."""
    if kind == "csv":
        return parse_csv(data, doc_id, vendor_id), ADAPTER_VERSIONS["csv"]
    if kind == "xlsx":
        return parse_xlsx(data, doc_id, vendor_id), ADAPTER_VERSIONS["xlsx"]
    if kind == "json":
        return parse_api(data, doc_id, vendor_id)
    raise IngestError(f"no structured adapter for {kind}")
